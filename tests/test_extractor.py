from io import BytesIO

import fitz
import pytest
from docx import Document

from engine.extractor import extract_text, extract_text_from_docx, extract_text_from_pdf, extract_text_from_txt


class UploadedFileStub(BytesIO):
    def __init__(self, name: str, data: bytes) -> None:
        super().__init__(data)
        self.name = name


def test_extract_text_from_txt_reads_utf8() -> None:
    file = BytesIO("This plays a crucial role.".encode("utf-8"))

    assert extract_text_from_txt(file) == "This plays a crucial role."


def test_extract_text_from_docx_reads_paragraphs() -> None:
    file = BytesIO()
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(file)

    assert extract_text_from_docx(file) == "First paragraph.\n\nSecond paragraph."


def test_extract_text_from_pdf_reads_selectable_text() -> None:
    file = BytesIO()
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Selectable PDF text.")
    document.save(file)
    document.close()

    assert "Selectable PDF text." in extract_text_from_pdf(file)


def test_extract_text_from_pdf_rejects_scanned_or_empty_pdf() -> None:
    file = BytesIO()
    document = fitz.open()
    document.new_page()
    document.save(file)
    document.close()

    with pytest.raises(ValueError, match="No extractable text"):
        extract_text_from_pdf(file)


def test_extract_text_routes_by_extension() -> None:
    uploaded_file = UploadedFileStub("sample.txt", b"Uploaded text.")

    assert extract_text(uploaded_file) == "Uploaded text."


def test_extract_text_rejects_unsupported_extensions() -> None:
    uploaded_file = UploadedFileStub("sample.rtf", b"Text")

    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(uploaded_file)
